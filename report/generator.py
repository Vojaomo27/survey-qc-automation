"""Generate a readable Excel QC report."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from core.models import RuleResult
from core.understanding import DataUnderstanding
from utils.constants import DEFAULT_REPORT_PREFIX, REPORT_SHEET_NAME


def generate_report(
    output_folder: str | Path,
    source_name: str,
    results: list[RuleResult],
    row_count: int,
    understanding: DataUnderstanding | None = None,
) -> Path:
    """Write summary and flagged records to a single formatted workbook."""
    folder = Path(output_folder)
    folder.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = folder / f"{DEFAULT_REPORT_PREFIX}_{timestamp}.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = REPORT_SHEET_NAME
    sheet.append(["Survey QC Report"])
    sheet.append(["Source", source_name])
    sheet.append(["Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    sheet.append(["Total records", row_count])
    sheet.append(["Total flagged", sum(result.flagged_count for result in results)])
    sheet.append([])
    sheet.append(["Rule", "Status", "Flagged", "Message"])
    for cell in sheet[7]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")

    for result in results:
        status = "Skipped" if result.skipped else "Complete"
        sheet.append([result.title, status, result.flagged_count, result.message])

    if understanding is not None:
        sheet.append([])
        sheet.append(["Key data metrics"])
        sheet.append(["Metric", "Value", "Potential impact"])
        for metric in understanding.metrics:
            sheet.append([metric.name, metric.value, metric.impact])
        sheet.append([])
        sheet.append(["Data issues and risks"])
        for issue in understanding.inconsistencies or ("No structural issues detected by the profiler.",):
            sheet.append([issue])

    for result in results:
        if not result.rows:
            continue
        sheet.append([])
        sheet.append([result.title])
        columns = result.columns or list(result.rows[0])
        sheet.append(columns)
        for row in result.rows:
            sheet.append([row.get(column) for column in columns])

    sheet.freeze_panes = "A8"
    for column in sheet.columns:
        width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 45)
        sheet.column_dimensions[column[0].column_letter].width = width
    workbook.save(path)
    return path